"""Prova única nas APIs reais Google Sheets/Drive, somente recursos e dados sintéticos.

Reaproveita SheetsProof (1e94527) e GoogleDrive. Falhas de rede marcadas como
"simuladas no cliente" são injetadas por este script; todo o resto é chamada real.
A credencial vem apenas de JANUS_GOOGLE_SA_JSON e nunca é impressa.
"""
import base64
import binascii
import hashlib
from io import BytesIO
import json
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from janus import extract, normalize
from drive_adapter import GoogleDrive
from experiments.sheets_poc.store import SheetsProof, Coordinator, Unconfirmed
from tests.test_sheets_poc import history_bytes

PROJECT = 'janus-homologacao'
ACCOUNT = 'janus-teste@janus-homologacao.iam.gserviceaccount.com'
SHEET = '1MaEfmY-7vnrS2c201evYopfoyRUDRXkzHg5udIFtKoo'
ROOT = '1z3FUCnOVMd4m0aSNE6cFBKrPDma0aQIw'
ENTRADA = '1BeIDpPK4_hbyjj8dWH6yjkblRnLSndeM'
PROCESSADOS = '1rIqKH_axohGIiITFd6emZ6C2OCteA8oN'
FALHAS = '17noWRSkyFEHUIlj07VeH0uE_mTd1bPq_'
CANONICAL = '12hFRwmYxrcpKa7X_NxHy3x0b3KpPhSWP' # Somente para confirmar ausência de acesso.
HISTORICAL = 'C000000000000000009' # Protocolo do histórico sintético de history_bytes().
SCOPES = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']


class Stop(Exception):
    pass


def credential_info():
    raw = os.environ.get('JANUS_GOOGLE_SA_JSON', '')
    if not raw.strip():
        raise Stop('Segredo ausente')
    try:
        info = json.loads(raw)
    except ValueError:
        try:
            info = json.loads(base64.b64decode(raw, validate=True))
        except (ValueError, binascii.Error):
            raise Stop('Segredo não é JSON nem JSON em base64') from None
    checks = {
        'type=service_account': info.get('type') == 'service_account',
        'project_id esperado': info.get('project_id') == PROJECT,
        'client_email esperado': info.get('client_email') == ACCOUNT,
        'private_key PEM': str(info.get('private_key', '')).startswith('-----BEGIN PRIVATE KEY-----'),
        'private_key_id presente': bool(info.get('private_key_id')),
        'token_uri Google': info.get('token_uri') == 'https://oauth2.googleapis.com/token',
    }
    return info, checks


class RealSheets:
    """Transporte real com a mesma interface de SimulatedSheets."""
    API = 'https://sheets.googleapis.com/v4/spreadsheets/' + SHEET

    def __init__(self, session):
        self.session, self.calls, self.fault = session, {'get':0, 'append':0}, None

    def get(self):
        self.calls['get'] += 1
        response = self.session.get(self.API + '/values/journal!A:F', timeout=60)
        response.raise_for_status()
        return response.json()

    def append(self, body, *, range, valueInputOption, insertDataOption):
        if (range,valueInputOption,insertDataOption) != ('journal!A:F','RAW','INSERT_ROWS'):
            raise ValueError('Somente append RAW INSERT_ROWS nesta prova')
        fault, self.fault = self.fault, None
        if fault == 'before':
            raise TimeoutError('Simulado no cliente: pedido não enviado')
        self.calls['append'] += 1
        response = self.session.post(self.API + '/values/' + range + ':append', timeout=60,
                                     params={'valueInputOption':valueInputOption,
                                             'insertDataOption':insertDataOption}, json=body)
        response.raise_for_status()
        if fault == 'after':
            raise TimeoutError('Simulado no cliente: resposta real descartada')
        return response.json()


class FlakyDrive:
    """Falha simulada no cliente antes do primeiro movimento real."""
    def __init__(self, drive):
        self.drive, self.failed = drive, False

    def archive(self, file_id, digest):
        if not self.failed:
            self.failed = True
            raise ConnectionError('Simulado no cliente: arquivamento interrompido')
        return self.drive.archive(file_id, digest)


def run(transport, drive, step):
    store = lambda: SheetsProof(transport, Coordinator()) # Cada chamada simula reinício a frio.

    start = store().state()
    step('Diário vazio antes da prova (garante execução única)', start['physical_rows'] == 0,
         f"{start['physical_rows']} linhas")
    if start['physical_rows']:
        raise Stop('Diário não vazio; prova já executada ou planilha alterada')

    files = {}
    with tempfile.TemporaryDirectory() as tmp:
        for item in drive.pending():
            name = drive._metadata(item.id)['name']
            if not name.startswith('SINTETICO_'):
                raise Stop('Arquivo não sintético na entrada')
            content = drive.read(item.id)
            path = Path(tmp) / 'f.pdf'; path.write_bytes(content)
            files[name] = (item.id, extract(path), hashlib.sha256(content).hexdigest())
    step('PDFs sintéticos lidos da entrada (Drive real)', len(files) == 4, ', '.join(sorted(files)))
    if len(files) != 4:
        raise Stop('Esperados 4 PDFs sintéticos')
    a, b, c, d = (files[f'SINTETICO_{x}.pdf'] for x in 'ABCD')
    step('Extração dos protocolos sintéticos', normalize(d[1].protocol) == HISTORICAL and
         len({normalize(x[1].protocol) for x in (a,b,c,d)}) == 4,
         ', '.join(x[1].protocol for x in (a,b,c,d)))

    history = history_bytes()
    step('Histórico sintético gravado (append real)', store().import_history(history) is True)
    rows = store().state()['physical_rows']
    step('Reimportar o mesmo histórico não acrescenta linhas', store().import_history(history) is False
         and store().state()['physical_rows'] == rows, f'{rows} linhas físicas')

    transport.fault = 'after'
    first = store().approve([a, b], approved=True)
    step('Lote 1 com resposta perdida (falha simulada no cliente após append real) confirmado por releitura',
         first['inserted'] == 2, json.dumps(first)[:120])
    calls, rows = transport.calls['append'], store().state()['physical_rows']
    again = store().approve([b, a], approved=True)
    step('Repetir lote 1 após reinício e ordem invertida: sem novo append',
         again == first and transport.calls['append'] == calls and store().state()['physical_rows'] == rows)

    transport.fault = 'before'
    try:
        store().approve([c, d], approved=True); unconfirmed = False
    except Unconfirmed:
        unconfirmed = True
    step('Lote 2 com timeout antes do envio (simulado no cliente) fica não confirmado',
         unconfirmed and store().state()['physical_rows'] == rows)
    second = store().approve([c, d], approved=True)
    step('Retomada do lote 2 após reinício; protocolo histórico ignorado',
         second['inserted'] == 1 and second['ignored'] == 1, json.dumps(second)[:120])

    state = store().state()
    late =[row for row in transport.get()['values'] if row[0] in ('R','C') and row[1] == first['batch_id']]
    transport.append({'values':[list(r) + ['']*(6-len(r)) for r in late]}, range='journal!A:F',
                     valueInputOption='RAW', insertDataOption='INSERT_ROWS')
    after = store().state()
    step('Cópia física tardia do lote 1 (append real duplicado) não cria registro lógico',
         after['physical_rows'] == state['physical_rows'] + len(late) and
         list(after['records']) == list(state['records']),
         f"{len(late)} linhas repetidas; {len(after['records'])} registros lógicos")

    def parents(item):
        return drive._metadata(item[0]).get('parents')
    step('Nenhum PDF arquivado antes da retomada', all(parents(x) == [ENTRADA] for x in (a,b,c,d)))
    resumed = store().resume(FlakyDrive(drive))
    step('Arquivamento com falha simulada no cliente preserva o PDF na entrada',
         resumed['archived'] == 2 and resumed['failures'] == [a[0]] and parents(a) == [ENTRADA],
         json.dumps({'archived':resumed['archived'],'failures':len(resumed['failures'])}))
    final = store().resume(drive)
    step('Retomada arquiva o restante (move real no Drive)', final == {'archived':1, 'failures':[]})
    step('Estado final no Drive: A,B,C em processados; D (ignorado) na entrada',
         all(parents(x) == [PROCESSADOS] for x in (a,b,c)) and parents(d) == [ENTRADA])
    step('Retomada repetida não move nem acrescenta', store().resume(drive) == {'archived':0, 'failures':[]})

    out = BytesIO(); store().export(out); out.seek(0)
    sheet = load_workbook(out, read_only=True)['JANUS 2026']
    protocols = [r[4] for r in sheet.iter_rows(min_row=3, values_only=True) if r and r[4]]
    step('Exportação .xlsx: histórico preservado + 3 novos, sem duplicata',
         protocols == [HISTORICAL, HISTORICAL] + [a[1].protocol, b[1].protocol, c[1].protocol],
         ', '.join(protocols))
    return store().state()['physical_rows']


def check_isolation(session, step):
    canonical = session.get('https://www.googleapis.com/drive/v3/files/' + CANONICAL, timeout=60,
                            params={'fields':'id'})
    step('Conta de serviço sem acesso à pasta canônica', canonical.status_code in (403, 404),
         f'HTTP {canonical.status_code}')
    visible, token = [], None
    while True:
        data = session.get('https://www.googleapis.com/drive/v3/files', timeout=60, params={
            'fields':'nextPageToken,files(id,parents)', 'pageSize':100, 'pageToken':token}).json()
        visible += data.get('files', [])
        token = data.get('nextPageToken')
        if not token:
            break
    allowed = {ROOT, ENTRADA, PROCESSADOS, FALHAS}
    step('Tudo que a conta de serviço vê está na raiz de teste',
         all(f['id'] == ROOT or set(f.get('parents', [])) <= allowed for f in visible),
         f'{len(visible)} itens visíveis')


def main():
    lines, ok = [], True
    def step(name, passed, detail=''):
        nonlocal ok
        ok = ok and bool(passed)
        lines.append(f"| {'OK' if passed else 'FALHOU'} | {name} | {detail} |")
        print(lines[-1], flush=True)
    error = None
    try:
        info, checks = credential_info()
        for name, passed in checks.items():
            step('Credencial: ' + name, passed)
        if not all(checks.values()):
            raise Stop('Estrutura da credencial inválida')
        from google.oauth2 import service_account
        from google.auth.transport.requests import AuthorizedSession
        session = AuthorizedSession(service_account.Credentials.from_service_account_info(info, scopes=SCOPES))
        del info
        transport = RealSheets(session)
        drive = GoogleDrive(session, ROOT, ENTRADA, PROCESSADOS, FALHAS)
        step('Estrutura de pastas de teste validada pelo adaptador (Drive real)', True)
        rows = run(transport, drive, step)
        check_isolation(session, step)
        step('Chamadas reais Sheets', True, f"values.get={transport.calls['get']}, values.append={transport.calls['append']}, linhas finais={rows}")
    except Exception as exc: # Somente tipo, status HTTP e mensagem da API; nunca a credencial.
        ok = False
        status = getattr(getattr(exc, 'response', None), 'status_code', '')
        message = ''
        try:
            message = exc.response.json().get('error', {}).get('message', '')[:200]
        except Exception:
            message = str(exc)[:200] if isinstance(exc, (Stop, ValueError, Unconfirmed)) else ''
        error = f'{type(exc).__name__} {status} {message}'.strip()
        print('ERRO:', error, flush=True)
    report = ['## ETAPA 3 — GATE GOOGLE REAL — RESULTADO DA EXECUÇÃO ÚNICA', '',
              f"Commit `{os.environ.get('GITHUB_SHA','local')[:7]}`, execução GitHub Actions "
              f"`{os.environ.get('GITHUB_RUN_ID','local')}`. Itens marcados *simulado no cliente* são falhas "
              'injetadas pelo script; todas as demais operações são chamadas reais às APIs Google.', '',
              '| Resultado | Verificação | Detalhe |', '|---|---|---|', *lines, '',
              f'**Erro:** {error}' if error else '',
              f"**Resultado automático:** {'todas as verificações OK' if ok else 'há falhas ou impedimento'}."]
    Path(os.environ.get('JANUS_REPORT', 'report.md')).write_text('\n'.join(report))
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
