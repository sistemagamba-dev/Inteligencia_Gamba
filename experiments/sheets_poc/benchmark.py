"""Somente dados inventados: append adversarial, reinício e dedup lógica."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
from time import perf_counter

from openpyxl import Workbook, load_workbook
from janus import HEADERS, extract
from drive_adapter import LocalDrive
from tests.fixtures import pdf, text
from experiments.sheets_poc.store import SheetsProof, Unconfirmed
from experiments.sheets_poc.fake import SimulatedSheets


def run():
    book=Workbook();sheet=book.active;sheet.title='JANUS 2026'
    sheet.append([]);sheet.append(HEADERS)
    historical=[]
    for i in range(3160):
        row=[i+1,None,'Pessoa histórica sintética',None,f'C{i % 3106:018}',
             'JANUS','Categoria sintética','Rua de teste 1','Pedido inventado','Bairro de teste']
        historical.append(row);sheet.append(row)
    source=BytesIO();book.save(source);book.close()
    remote=SimulatedSheets();client=SheetsProof(remote)
    start=perf_counter();client.import_history(source.getvalue())
    with tempfile.TemporaryDirectory(prefix='janus-sheets-append-') as tmp:
        drive=LocalDrive(Path(tmp)/'drive');items=[]
        for i in range(100):
            protocol=f'C{0 if i==0 else 10001 if i==2 else 10000+i:018}'
            path=drive.root/'entrada'/f'{i:03}.pdf'
            pdf(path,text(protocol,events=bool(i%2)),pages=4 if i%2 else 2)
            items.append((path.name,extract(path),hashlib.sha256(path.read_bytes()).hexdigest()))
        ready=perf_counter()
        assert not client.state()['records']
        remote.failures=['deferred']
        try:
            client.approve(items[:50],approved=True)
            raise AssertionError('Primeiro lote deveria estar indeterminado')
        except Unconfirmed:
            pass
        # Reinício completo do cliente, sem memória de reservas/IDs do anterior.
        fresh=SheetsProof(remote)
        second=fresh.approve(items[50:],approved=True)
        assert second['inserted']==50
        prefix=[row.copy() for row in remote.rows]
        remote.deliver_deferred(copies=2) # Duas chegadas tardias do primeiro pedido.
        assert remote.rows[:len(prefix)]==prefix
        first=SheetsProof(remote).approve(items[:50],approved=True)
        assert (first['inserted'],first['ignored'])==(48,2)
        committed=perf_counter()
        state=fresh.state();count_before=len(state['records']);commits_before=remote.commits
        replay=SheetsProof(remote).approve(items[:50],approved=True)
        assert replay==first and remote.commits==commits_before
        count_after=len(fresh.state()['records'])
        reprocess=SheetsProof(remote).approve(items,approved=True)
        assert (reprocess['inserted'],reprocess['ignored'])==(0,100)
        archive=SheetsProof(remote).resume(drive)
        assert archive=={'archived':98,'failures':[]}
        output=BytesIO();client.export(output);output.seek(0)
        exported=load_workbook(output,read_only=True)
        rows=list(exported.active.iter_rows(min_row=3,max_col=10,values_only=True))
        assert [list(r) for r in rows[:3160]]==historical
        assert len(rows)==3258
        assert rows[3160][0]==3161
        assert len({r[4] for r in rows[3160:]})==98
        assert len(list((drive.root/'entrada').glob('*.pdf')))==2
        assert len(list((drive.root/'processados').glob('*.pdf')))==98
        exported.close()
        physical_records=sum(row[0]=='R' for row in remote.rows)
        return {'backend':'Sheets simulado; values.append RAW INSERT_ROWS; sem Google real',
                'base_commit':'15c3b49aac3dbee1bbdf7dee8fb77b1fe2270488',
                'synthetic_history':{'rows':3160,'unique_protocols':3106,'duplicate_excess':54},
                'synthetic_pdfs':100,'pages':300,'logical_records_before_approval':0,
                'second_batch_confirmed_before_delayed_first':second,
                'first_batch_after_delayed_double_delivery':first,
                'physical_new_record_events':physical_records,'logical_new_records':98,
                'physical_repetitions_consolidated':physical_records-98,
                'additional_logical_records_on_replay':count_after-count_before,
                'additional_appends_on_replay':0,
                'full_reprocessing':reprocess,'archive':archive,
                'preserved_pending_pdfs':2,'export_rows':len(rows),
                'historical_values_identical':True,'confirmed_prefix_unchanged':True,
                'append_calls':remote.append_calls,'applied_deliveries':remote.commits,
                'largest_serialized_append_bytes':remote.max_bytes,
                'seconds':{'import_generate_extract':round(ready-start,6),
                           'adversarial_approval_and_reconciliation':round(committed-ready,6),
                           'total_including_export_archive':round(perf_counter()-start,6)},
                'known_data_integrity_failures_in_executed_scenarios':0,
                'google_real_validated':False}


if __name__=='__main__':
    print(json.dumps(run(),ensure_ascii=False,indent=2))
