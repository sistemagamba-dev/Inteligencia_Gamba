"""Lote inteiramente sintético: gera, extrai, grava, repete e exporta 100 PDFs."""
import json
from pathlib import Path
import tempfile
from time import perf_counter
from openpyxl import load_workbook
from janus import History, TrialStore, extract
from tests.fixtures import pdf, text


def run():
    with tempfile.TemporaryDirectory(prefix='gamba-batch-') as tmp:
        root = Path(tmp)
        start = perf_counter()
        for n in range(100):
            pdf(root/f'{n:03}.pdf', text(f'C{n:018}', events=bool(n % 2)), pages=4 if n % 2 else 2)
        generated = perf_counter()
        items = [(p.stem, extract(p)) for p in sorted(root.glob('*.pdf'))]
        assert all(not r.warnings for _, r in items)
        extracted = perf_counter()
        store = TrialStore(root/'test.sqlite', History([], set(), 1, 0))
        moved = set()
        first = store.approve(items, approved=True, archive=moved.add)
        committed = perf_counter()
        repeat = store.approve(items, approved=True, archive=moved.add)
        repeated = perf_counter()
        store.export(root/'export.xlsx')
        w = load_workbook(root/'export.xlsx', read_only=True)
        assert len(list(w.active.iter_rows(min_row=3))) == 100
        w.close()
        assert first == {'inserted': 100, 'ignored': 0}
        assert repeat == {'inserted': 0, 'ignored': 100}
        assert len(moved) == 100
        store.close()
        return {'pdfs': 100, 'pages': 300, 'layouts': {'Zeladoria': 50, 'Eventos': 50},
                'first': first, 'repeat': repeat, 'unique_archive_operations': len(moved),
                'seconds': {'generation': round(generated-start, 6),
                            'extraction': round(extracted-generated, 6),
                            'first_save_and_archive_simulation': round(committed-extracted, 6),
                            'repeated_batch': round(repeated-committed, 6),
                            'total_including_export_check': round(perf_counter()-start, 6)}}


if __name__ == '__main__':
    print(json.dumps(run(), indent=2))
