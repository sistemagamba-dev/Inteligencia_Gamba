"""Operações da equipe técnica, nunca exigidas dos assessores."""
import argparse
from io import BytesIO
from pathlib import Path
import os
import sqlite3
import tempfile
from types import SimpleNamespace
import uuid

from janus import HEADERS
from openpyxl import Workbook, load_workbook
from storage import Store

EXCEL_CELL = 32767


def restore(database, backup):
    database, backup = Path(database).resolve(), Path(backup).resolve()
    if database == backup:
        raise ValueError('Origem e destino devem ser diferentes')
    # SQLite abre só em leitura. Não cria um backup ausente por engano.
    source = sqlite3.connect(backup.as_uri() + '?mode=ro', uri=True)
    try:
        if source.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
            raise ValueError('Cópia inválida')
        expected = {'records', 'outbox', 'history', 'meta', 'batches'}
        if not expected.issubset({r[0] for r in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}):
            raise ValueError('Cópia incompatível')
        store = Store(database)
        try:
            with store.exclusive():
                # Aplicativo deve estar parado: nenhuma conexão pode sobreviver à troca.
                preserved = database.with_name(database.name + '.before-restore-' + uuid.uuid4().hex + '.bak')
                target = sqlite3.connect(preserved)
                store.db.backup(target)
                target.close()
                with tempfile.NamedTemporaryFile(dir=database.parent, delete=False) as tmp:
                    stage = Path(tmp.name)
                try:
                    destination = sqlite3.connect(stage)
                    source.backup(destination)
                    destination.close()
                    store.close()
                    os.replace(stage, database)
                finally:
                    stage.unlink(missing_ok=True)
        finally:
            store.close()
    finally:
        source.close()


def history_metrics(source):
    """Somente métricas agregadas do histórico; nenhum valor de célula."""
    from experiments.sheets_poc.store import history_plan, measure, MAX_BYTES, MAX_CELL
    h, _, events, manifest, chunks = history_plan(source)
    size, cell = measure(events + [manifest])
    return {'linhas': len(h.rows), 'com_protocolo': h.protocol_rows, 'unicos': len(h.keys),
            'excesso_duplicados': h.protocol_rows - len(h.keys), 'bytes_json': size,
            'maior_celula': cell, 'appends': len(chunks), 'cabe_em_um_append': size <= MAX_BYTES,
            'celulas_ok': cell <= MAX_CELL}


def import_history_sheets(path, components):
    """Valida o arquivo antes de abrir credenciais; components só é chamado se passar."""
    source = Path(path).read_bytes()
    if not history_metrics(source)['celulas_ok']:
        raise ValueError('Histórico tem célula acima de 50 mil caracteres; nada foi gravado')
    return components().proof.import_history(source)


def backup_sheets(path, components):
    """Cópia local nova: base lógica (se íntegra) + aba journal com A:F brutos. Só lê o Google."""
    from experiments.sheets_poc.store import SheetsProof
    path = Path(path)
    if path.exists():
        raise ValueError('Destino já existe; escolha um arquivo novo')
    response = components().proof.transport.get() # Uma leitura: as duas abas vêm do mesmo instante.
    rows = [list(r) + [''] * (6-len(r)) for r in response.get('values', [])]
    if any(len(str(v)) > EXCEL_CELL for row in rows for v in row):
        raise ValueError('Célula acima do limite do Excel; cópia não gerada')
    logical, buffer = True, BytesIO()
    try:
        SheetsProof(SimpleNamespace(get=lambda: response)).export(buffer)
        book = load_workbook(BytesIO(buffer.getvalue()))
    except ValueError:
        logical, book = False, Workbook()  # Diário inconsistente: preserva ao menos a cópia bruta.
        book.remove(book.active)
    sheet = book.create_sheet('journal')
    for row in rows:
        sheet.append(row)
    for row in (r for ws in book.worksheets for r in ws.iter_rows()):
        for cell in row:
            if cell.data_type == 'f': cell.data_type = 's'
    output = BytesIO(); book.save(output); book.close()
    with open(path, 'xb') as file:
        file.write(output.getvalue())
    return {'linhas_journal': len(rows), 'base_logica': logical}


def demo(root):
    root = Path(root)
    if (root/'janus.sqlite').exists():
        raise ValueError('Use uma pasta nova para a demonstração')
    history = demo_files(root)
    store = Store(root/'janus.sqlite')
    store.import_history(history)
    store.close()


def demo_files(root):
    """PDFs e histórico inventados; devolve o caminho do histórico sintético."""
    from drive_adapter import LocalDrive
    from tests.fixtures import pdf, text, PROTOCOL
    root = Path(root)
    drive = LocalDrive(root/'drive')
    if any((drive.root/'entrada').iterdir()):
        raise ValueError('Use uma pasta nova para a demonstração')
    history = root/'historico-sintetico.xlsx'
    book = Workbook()
    book.active.title = 'JANUS 2026'
    book.active.append([])
    book.active.append(HEADERS)
    for order in (4, 9):
        book.active.append([order, None, 'Pessoa histórica sintética', None, PROTOCOL, 'JANUS'])
    book.save(history)
    book.close()
    for name, protocol, events, missing in [
        ('01-zeladoria', 'C000000000000000001', False, False),
        ('02-eventos', 'C000000000000000002', True, False),
        ('03-a-conferir', 'C000000000000000003', False, True),
        ('04-duplicado', PROTOCOL, False, False)]:
        content = text(protocol, events=events)
        if missing:
            content = content.replace('Enquadramento: Manutencao sintetica\n', '')
        pdf(drive.root/'entrada'/f'{name}.pdf', content, date=None if missing else 'D:20261028010000Z')
    (drive.root/'entrada'/'05-invalido.pdf').write_bytes(b'PDF sintetico invalido')
    return history


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    cmd = sub.add_parser('demo'); cmd.add_argument('root', type=Path)
    cmd = sub.add_parser('import-history-sheets', help='Importa uma única vez o histórico no diário do Sheets')
    cmd.add_argument('file', type=Path)
    cmd = sub.add_parser('verificar-historico', help='Somente leitura: métricas agregadas do histórico')
    cmd.add_argument('file', type=Path)
    cmd = sub.add_parser('backup-sheets', help='Somente leitura: copia base e diário para um .xlsx novo')
    cmd.add_argument('file', type=Path)
    for name in ('import-history', 'backup', 'restore'):
        cmd = sub.add_parser(name)
        cmd.add_argument('database', type=Path)
        cmd.add_argument('file', type=Path)
    args = parser.parse_args()
    if args.command == 'demo':
        demo(args.root)
    elif args.command == 'verificar-historico':
        for key, value in history_metrics(args.file.read_bytes()).items():
            print(f'{key}: {value}')
    elif args.command == 'import-history-sheets':
        from runtime import components
        imported = import_history_sheets(args.file, components)
        print('Histórico importado.' if imported else 'Mesmo histórico já importado; nada alterado.')
    elif args.command == 'backup-sheets':
        from runtime import components
        result = backup_sheets(args.file, components)
        print(f"Cópia salva: {result['linhas_journal']} linhas do diário." +
              ('' if result['base_logica'] else ' ATENÇÃO: diário inconsistente; só a cópia bruta foi salva.'))
    elif args.command == 'restore':
        restore(args.database, args.file)
    else:
        store = Store(args.database)
        try:
            if args.command == 'backup':
                store.backup(args.file)
            else:
                store.import_history(args.file)
        finally:
            store.close()
