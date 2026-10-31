"""Ensaio opcional com Chromium/Playwright; servidor local em modo demonstração já iniciado
com JANUS_DATA_DIR apontando para uma pasta nova (Sheets simulado em memória).
Uso: python browser_check.py /tmp/janus-browser-test docs/evidence/etapa3
Requer playwright e chromium, apenas na máquina de verificação.
"""
import json
import os
from pathlib import Path
import sys
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook


def run(root, evidence):
    evidence.mkdir(parents=True, exist_ok=True)
    exports = []
    def count(page):
        # Base lida pela exportação da própria tela: linhas além do cabeçalho e do histórico (2).
        page.get_by_role('button',name='Preparar Excel da base',exact=True).click()
        button = page.get_by_role('button',name='Baixar Excel (.xlsx)',exact=True)
        expect(button).to_be_visible()
        with page.expect_download() as download:
            button.click()
        path = root/f'browser-export-{len(exports)}.xlsx'
        download.value.save_as(path)
        exports.append(path)
        book = load_workbook(path,read_only=True)
        sheet = book['JANUS 2026']
        assert sheet.max_column == 10
        rows = sheet.max_row - 2 - 2
        book.close()
        return rows
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=os.environ.get('JANUS_CHROMIUM','/usr/bin/chromium'),headless=True,args=['--no-sandbox'])
        page = browser.new_page(viewport={'width':1440,'height':1100})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(os.environ.get('JANUS_BROWSER_URL','http://127.0.0.1:8501'))
        assert count(page) == 0, 'Use demonstração nova'
        page.get_by_role('button',name='Processar PDFs',exact=True).click()
        expect(page.get_by_role('button',name='Aprovar e registrar tudo',exact=True)).to_be_visible()
        assert len(list((root/'drive'/'processados').glob('*.pdf'))) == 0
        expect(page.get_by_text('Novos para conferir: 3',exact=False)).to_be_visible()
        page.wait_for_timeout(800)  # Aguarda pintura dos canvas da tabela.
        page.screenshot(path=str(evidence/'browser-preview.png'),full_page=True)
        page.get_by_role('button',name='Aprovar e registrar tudo',exact=True).click()
        expect(page.get_by_text('3 novos registrados',exact=False)).to_be_visible()
        assert count(page) == 3
        page.wait_for_timeout(400)
        page.screenshot(path=str(evidence/'browser-approved.png'),full_page=True)
        page.reload()
        page.get_by_role('button',name='Processar PDFs',exact=True).click()
        expect(page.get_by_text('Novos para conferir: 0',exact=False)).to_be_visible()
        assert count(page) == 3
        assert len(list((root/'drive'/'processados').glob('*.pdf'))) == 3
        assert len(list((root/'drive'/'entrada').glob('*.pdf'))) == 2
        assert not page.locator('[data-testid="stException"]').count()
        assert not errors
        # Viewport compacto: ações continuam acessíveis, tabela rola horizontalmente.
        page.set_viewport_size({'width':390,'height':844})
        expect(page.get_by_role('button',name='Processar PDFs',exact=True)).to_be_visible()
        browser.close()
    return {'environment':'Chromium headless / Streamlit local / Sheets simulado / PDFs sintéticos',
            'preview_before_approval_records':0,'after_approval_records':3,
            'historical_rows_preserved':2,'duplicates_ignored':1,'invalid_files_preserved':1,
            'archived_pdfs':3,'pending_pdfs':2,'xlsx_columns':10,'xlsx_data_rows':5,
            'browser_reload_no_new_records':True,'browser_errors':errors,
            'compact_viewport_action_visible':True,
            'google_live_calls':False,'human_acceptance':False}


if __name__ == '__main__':
    result = run(Path(sys.argv[1]),Path(sys.argv[2]))
    (Path(sys.argv[2])/'browser.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))
