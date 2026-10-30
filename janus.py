"""Etapa 1: leitura local, sem integrações operacionais nem logging de documentos."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import re
import sqlite3
from zoneinfo import ZoneInfo

from openpyxl import load_workbook, Workbook
from pypdf import PdfReader

HEADERS = ('ordem', 'PREENCHIDO EM', 'Solicitante', 'Data do Envio',
           'PROTOCOLO', 'SISTEMA', 'ENQUADRAMENTO', 'ENDEREÇO',
           'OCORRÊNCIA- Descrição sintética do pedido', 'BAIRRO')


def normalize(value: object) -> str:
    return re.sub(r'[^A-Z0-9]', '', str(value or '').upper())


def valid_protocol(value: object) -> bool:
    return bool(re.fullmatch(r'C\d{18}', normalize(value)))


def creation_time(raw: str | None) -> datetime | None:
    """Não presume fuso para metadados sem offset explícito."""
    if not isinstance(raw, str) or not raw:
        return None
    m = re.fullmatch(r"D:(\d{14})(Z|[+-]\d{2}'?\d{2}'?)", raw)
    if not m:
        return None
    try:
        zone = m[2].replace("'", '')
        offset = timedelta() if zone == 'Z' else timedelta(
            minutes=(int(zone[1:3]) * 60 + int(zone[3:5])) * (1 if zone[0] == '+' else -1))
        return datetime.strptime(m[1], '%Y%m%d%H%M%S').replace(
            tzinfo=timezone(offset)).astimezone(ZoneInfo('America/Sao_Paulo'))
    except ValueError:
        return None


def compact(value: str) -> str:
    return re.sub(r'\s+', ' ', value).strip()


def capture(text: str, start: str, end: str) -> str:
    match = re.search(start + r'\s*(.*?)' + end, text, re.S | re.I)
    return compact(match[1]) if match else ''


@dataclass
class Record:
    protocol: str = ''
    kind: str = ''
    requester: str = ''
    created: str = ''
    category: str = ''
    address: str = ''
    neighborhood: str = ''
    description: str = ''
    full_address: str = ''
    warnings: list[str] = field(default_factory=list)

    def row(self, sequence: int) -> list:
        return [sequence, None, self.requester or None, self.created or None,
                self.protocol, 'JANUS', self.category or None, self.address or None,
                self.description or None, self.neighborhood or None]


def parse_text(text: str, raw_date: str | None) -> Record:
    keys = {normalize(v) for v in re.findall(r'^Protocolo:\s*([^\n]+)', text, re.M | re.I)}
    # Mantém rodapés como delimitadores para não incorporar anexos ao pedido.
    r = Record()
    header = re.search(r'^Protocolo:\s*([^\n]+)', text, re.M | re.I)
    r.protocol = compact(header[1]) if header else ''
    if not valid_protocol(r.protocol) or keys != {normalize(r.protocol)}:
        r.warnings.append('protocolo_invalido_ou_conflitante')
        r.protocol = ''
    r.kind = capture(text, r'Serviço:', r'\n')
    r.requester = capture(text, r'Responsável pelo\s+preenchimento:',
                          r'\n(?:E-mail|Tipo de pedido)')
    if 'Zeladoria' in r.kind:
        r.category = capture(text, r'Enquadramento:', r'\nLocal:')
        r.full_address = capture(text, r'Endereço:', r'\nDescrição sintética')
        r.description = capture(text, r'Descrição sintética do\s+pedido:',
                                r'(?:\nSolicitação|\nProtocolo:|\Z)')
    elif 'Eventos' in r.kind:
        r.category = capture(text, r'\nEnquadramento\s*\nEnquadramento\s+', r'\n')
        r.full_address = capture(text, r'Endereço do evento\s+', r'\nHorário de início')
        r.description = capture(text, r'Descrição objetiva do\s+pedido\s*\n',
                                r'\n(?:Informações\s+complementares|Solicitação|Protocolo:)|\Z')
    else:
        r.warnings.append('tipo_nao_suportado_revisar')
    parts = [compact(p) for p in r.full_address.split(',') if compact(p)]
    # Endereço completo fica preservado separadamente para revisão, inclusive CEP.
    if len(parts) >= 3 and re.search(r'Santo André\s*-\s*SP', parts[-2] if len(parts) > 3 else parts[-1], re.I):
        r.address, r.neighborhood = parts[0], parts[1]
    else:
        r.address = r.full_address
        if r.full_address:
            r.warnings.append('bairro_ambiguo_revisar')
    created = creation_time(raw_date)
    r.created = created.date().isoformat() if created else ''
    for name in ['requester', 'created', 'category', 'address', 'neighborhood', 'description']:
        if not getattr(r, name):
            r.warnings.append('ausente_' + name)
    return r


def extract(path: Path) -> Record:
    reader = PdfReader(path)
    return parse_text('\n'.join(page.extract_text() or '' for page in reader.pages),
                      (reader.metadata or {}).get('/CreationDate'))


@dataclass
class History:
    rows: list[list]
    keys: set[str]
    next_order: int
    protocol_rows: int


def read_history(path: Path) -> History:
    """Valores A:J, sem salvar a origem. Não tenta sanear duplicatas históricas."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book['JANUS 2026']
        headers = tuple(next(sheet.iter_rows(min_row=2, max_row=2, max_col=10, values_only=True), ()))
        if headers != HEADERS:
            raise ValueError('Cabeçalhos A:J divergentes; revisar mapeamento')
        rows, keys, orders, count = [], set(), [0], 0
        for index, values in enumerate(sheet.iter_rows(min_row=3, max_col=10, values_only=True), 3):
            rows.append(list(values))
            if values[4] is not None and str(values[4]).strip():
                count += 1
                keys.add(normalize(values[4]))
            if isinstance(values[0], (int, float)) and float(values[0]).is_integer():
                orders.append(int(values[0]))
        return History(rows, keys, max(orders) + 1, count)
    finally:
        book.close()


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class TrialStore:
    """Simulador local SQLite: commit + outbox recuperável; sem acesso ao Drive.

    archive(file_id) deve ser idempotente. Falha após mover e antes de marcar
    exige que o adaptador futuro confirme o destino ao repetir a operação.
    """
    def __init__(self, path: Path, history: History):
        self.db = sqlite3.connect(path)
        self.history = history
        self.db.executescript('''
            CREATE TABLE IF NOT EXISTS records(key TEXT PRIMARY KEY, payload TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS outbox(file_id TEXT PRIMARY KEY, key TEXT NOT NULL,
                                             archived INTEGER NOT NULL DEFAULT 0);
        ''')

    def approve(self, items: list[tuple[str, Record]], *, approved: bool,
                archive, after_commit=None) -> dict:
        if not approved:
            raise ValueError('Aprovação humana obrigatória')
        if any(not valid_protocol(r.protocol) for _, r in items):
            raise ValueError('Lote contém protocolo sem identificação confiável')
        inserted = ignored = 0
        with self.db:
            # Serializa verificação de duplicidade e sequência entre escritores.
            self.db.execute('BEGIN IMMEDIATE')
            for file_id, record in items:
                key = normalize(record.protocol)
                exists = key in self.history.keys or self.db.execute(
                    'SELECT 1 FROM records WHERE key=?', (key,)).fetchone()
                if exists:
                    ignored += 1
                else:
                    order = self.history.next_order + self.db.execute('SELECT count(*) FROM records').fetchone()[0]
                    self.db.execute('INSERT INTO records VALUES (?,?)',
                                    (key, json.dumps(record.row(order), ensure_ascii=False)))
                    inserted += 1
                bound = self.db.execute('SELECT key FROM outbox WHERE file_id=?', (file_id,)).fetchone()
                if bound and bound[0] != key:
                    raise ValueError('Arquivo já vinculado a outro protocolo')
                self.db.execute('INSERT OR IGNORE INTO outbox(file_id,key) VALUES (?,?)', (file_id, key))
        if after_commit:
            after_commit()
        self.resume(archive)
        return {'inserted': inserted, 'ignored': ignored}

    def resume(self, archive):
        for (file_id,) in self.db.execute('SELECT file_id FROM outbox WHERE archived=0').fetchall():
            archive(file_id)
            with self.db:
                self.db.execute('UPDATE outbox SET archived=1 WHERE file_id=?', (file_id,))

    def export(self, destination: Path):
        book = Workbook()
        sheet = book.active
        sheet.title = 'JANUS 2026'
        sheet.append([])
        sheet.append(HEADERS)
        for row in self.history.rows:
            sheet.append(row)
        for (payload,) in self.db.execute('SELECT payload FROM records ORDER BY rowid'):
            sheet.append(json.loads(payload))
        # Impede fórmulas injetadas sem alterar os valores literais.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.data_type == 'f':
                    cell.data_type = 's'
        book.save(destination)
        book.close()

    def close(self):
        self.db.close()
