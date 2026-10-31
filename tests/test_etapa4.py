"""Etapa 4: ambiente explícito, histórico em partes e cópia de segurança. Somente fakes e dados inventados."""
from io import BytesIO
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
import streamlit as st
from streamlit.testing.v1 import AppTest
from drive_adapter import LocalDrive
from experiments.sheets_poc.fake import SimulatedSheets
from experiments.sheets_poc.store import SheetsProof, Unconfirmed, MAX_BYTES
from janus import HEADERS, read_history
from manage import backup_sheets, demo_files, history_metrics, import_history_sheets
import runtime
from workflow import SheetsWorkflow

SHEET, ROOT = runtime.HOMOLOGACAO['sheet'], runtime.HOMOLOGACAO['root']
BASE = {'JANUS_SHEET_ID': 'planilha_nova', 'JANUS_ENTRADA': 'inbox', 'JANUS_PROCESSADOS': 'done',
        'JANUS_FALHAS': 'failed'}


def env(**values):
    clean = {k: v for k, v in os.environ.items() if not k.startswith('JANUS_')}
    return patch.dict(os.environ, {**clean, **BASE, **values}, clear=True)


def session(root_name):
    response = lambda data: Mock(json=Mock(return_value=data))
    s = Mock()
    s.request.side_effect = ([response({'name': root_name, 'mimeType': 'application/vnd.google-apps.folder'})] +
                             [response({'parents': ['raiz'], 'mimeType': 'application/vnd.google-apps.folder'})]*3)
    return s


def history_file(path, rows):
    book = Workbook(); sheet = book.active; sheet.title = 'JANUS 2026'
    sheet.append([]); sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    book.save(path); book.close()
    return path


class EnvironmentTests(unittest.TestCase):
    def build(self, ambiente, name):
        with env(JANUS_AMBIENTE=ambiente, JANUS_TEST_ROOT='raiz'):
            return runtime.google_components(session(name), runtime.google_settings())

    def test_root_prefix_matrix(self):
        for ambiente, name, ok in [('homologacao', 'JANUS_TESTE_HOMOLOGACAO', True),
                                   ('homologacao', 'JANUS_PRODUCAO_GABINETE', False),
                                   ('producao', 'JANUS_PRODUCAO_GABINETE', True),
                                   ('producao', 'JANUS_TESTE_HOMOLOGACAO', False),
                                   ('producao', 'JANUS_TESTE_JANUS_PRODUCAO', False),
                                   ('producao', 'Operacional', False),
                                   ('homologacao', 'Operacional', False)]:
            with self.subTest(ambiente=ambiente, name=name):
                if ok:
                    self.assertTrue(self.build(ambiente, name).drive.entrada)
                else:
                    with self.assertRaisesRegex(ValueError, 'raiz'):
                        self.build(ambiente, name)

    def test_default_is_homologacao_and_unknown_refuses(self):
        with env():
            self.assertEqual((runtime.ambiente(), runtime.producao()), ('homologacao', False))
        for value in ('produção', 'PRODUCAO', 'teste', ''):
            with self.subTest(value=value), env(JANUS_AMBIENTE=value):
                with self.assertRaisesRegex(ValueError, 'Ambiente desconhecido'):
                    runtime.components()

    def test_production_refuses_homologation_resources_before_credentials(self):
        for extra in ({'JANUS_SHEET_ID': SHEET, 'JANUS_TEST_ROOT': 'raiz'}, {'JANUS_TEST_ROOT': ROOT},
                      {'JANUS_ROOT': ROOT}):
            with self.subTest(extra=extra), env(JANUS_AMBIENTE='producao', JANUS_MODE='google', **extra):
                with self.assertRaisesRegex(ValueError, 'homologação'):
                    runtime.components() # Sem JANUS_GOOGLE_SA_JSON: falharia com KeyError se lesse credenciais.
        with env(JANUS_SHEET_ID=SHEET, JANUS_TEST_ROOT=ROOT):
            self.assertEqual(runtime.google_settings()['folders'][0], ROOT) # Homologação inalterada.

    def test_production_requires_google_mode(self):
        with env(JANUS_AMBIENTE='producao'):
            with self.assertRaisesRegex(ValueError, 'JANUS_MODE=google'):
                runtime.components()

    def test_root_alias(self):
        with env(JANUS_ROOT='raiz'):
            self.assertEqual(runtime.google_settings()['folders'], ['raiz', 'inbox', 'done', 'failed'])
        with env(JANUS_ROOT='raiz', JANUS_TEST_ROOT='raiz'):
            self.assertEqual(runtime.google_settings()['prefix'], 'JANUS_TESTE')
        with env(JANUS_ROOT='raiz', JANUS_TEST_ROOT='outra'):
            with self.assertRaisesRegex(ValueError, 'divergentes'):
                runtime.google_settings()
        with env():
            with self.assertRaisesRegex(ValueError, 'JANUS_ROOT'):
                runtime.google_settings()


class CaptionTests(unittest.TestCase):
    def setUp(self):
        st.cache_resource.clear()
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        proof = SheetsProof(SimulatedSheets())
        proof.import_history(demo_files(root).read_bytes())
        patcher = patch.object(runtime, 'components', return_value=SheetsWorkflow(proof, LocalDrive(root/'drive')))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(st.cache_resource.clear)

    def run_app(self, ambiente):
        with env(JANUS_AMBIENTE=ambiente):
            app = AppTest.from_file('app.py', default_timeout=30).run()
        self.assertFalse(app.exception)
        return app

    def test_caption_per_environment(self):
        app = self.run_app('homologacao')
        self.assertEqual(app.caption[0].value, 'INTELIGÊNCIA GAMBÁ  /  AMBIENTE DE TESTE')
        self.assertIn('Demonstração', app.info[0].value)
        app = self.run_app('producao')
        self.assertEqual(app.caption[0].value, 'INTELIGÊNCIA GAMBÁ  /  OPERAÇÃO')
        texts = [e.value for e in (*app.caption, *app.info, *app.markdown, *app.error)]
        self.assertFalse([t for t in texts if 'TESTE' in t.upper() or 'Demonstração' in t])

    def test_unknown_environment_loads_nothing(self):
        app = self.run_app('outro')
        self.assertIn('equipe técnica', app.error[0].value)
        self.assertFalse(app.caption)
        runtime.components.assert_not_called()


class HistoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def big(self, name='grande.xlsx', rows=300):
        # ~20 kB por linha: ~6 MB de corpo, forçando ao menos 3 partes.
        return history_file(self.root/name, [[i, None, f'Pessoa {i}', None, f'C{i:018}', 'x'*20000]
                                             for i in range(1, rows+1)] + [[None]*10])

    def test_metrics_are_aggregate_only(self):
        history = demo_files(self.root)
        metrics = history_metrics(history.read_bytes())
        self.assertEqual({k: metrics[k] for k in ('linhas', 'com_protocolo', 'unicos', 'excesso_duplicados',
                                                  'appends', 'cabe_em_um_append', 'celulas_ok')},
                         {'linhas': 2, 'com_protocolo': 2, 'unicos': 1, 'excesso_duplicados': 1, 'appends': 1,
                          'cabe_em_um_append': True, 'celulas_ok': True})
        output = subprocess.run([sys.executable, 'manage.py', 'verificar-historico', str(history)],
                                capture_output=True, text=True, check=True).stdout
        self.assertIn('excesso_duplicados: 1', output)
        self.assertNotIn('Pessoa', output)
        self.assertNotIn('C1234', output)

    def test_big_history_imports_in_chunks_and_reads_back(self):
        path = self.big()
        metrics = history_metrics(path.read_bytes())
        self.assertFalse(metrics['cabe_em_um_append'])
        self.assertGreaterEqual(metrics['appends'], 4) # >= 3 partes H + manifesto I.
        remote = SimulatedSheets()
        self.assertTrue(SheetsProof(remote).import_history(path.read_bytes()))
        self.assertEqual(remote.append_calls, metrics['appends'])
        self.assertLessEqual(remote.max_bytes, MAX_BYTES)
        self.assertEqual(SheetsProof(remote).state()['history'], read_history(path).rows)
        self.assertFalse(SheetsProof(remote).import_history(path.read_bytes()))
        with self.assertRaisesRegex(ValueError, 'Não substituir'):
            SheetsProof(remote).import_history(history_file(self.root/'outro.xlsx', [[1]]).read_bytes())

    def test_interrupted_chunked_import_is_retried_once(self):
        path = self.big()
        for failure in ('before', 'after', 'deferred'):
            with self.subTest(failure=failure):
                remote = SimulatedSheets()
                remote.failures = [None, failure]
                with self.assertRaises(Unconfirmed):
                    SheetsProof(remote).import_history(path.read_bytes())
                remote.deliver_deferred() # Parte em trânsito chega antes da nova tentativa.
                self.assertIsNone(SheetsProof(remote).state()['history_source'])
                self.assertTrue(SheetsProof(remote).import_history(path.read_bytes()))
                state = SheetsProof(remote).state()
                self.assertEqual(state['history'], read_history(path).rows)
                self.assertEqual(sum(r[0] == 'I' for r in remote.rows), 1)

    def test_chunked_import_refused_when_journal_has_records(self):
        remote = SimulatedSheets()
        remote.rows = [['R', 'lote', 'C000000000000000001', 'JANUS:C000000000000000001', '[]', '{}']]
        with self.assertRaisesRegex(ValueError, 'registros'):
            SheetsProof(remote).import_history(self.big(rows=120).read_bytes())
        self.assertEqual(remote.append_calls, 0)

    def test_cell_over_limit_refused_before_credentials(self):
        # Cada linha vira uma célula JSON no diário: duas células longas passam de 50 mil.
        path = history_file(self.root/'celula.xlsx', [[1, None, 'y'*30000, None, 'C000000000000000001', 'x'*30000]])
        self.assertFalse(history_metrics(path.read_bytes())['celulas_ok'])
        components = Mock()
        with self.assertRaisesRegex(ValueError, '50 mil'):
            import_history_sheets(path, components)
        components.assert_not_called()


class BackupTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        self.remote = SimulatedSheets()
        self.flow = SheetsWorkflow(SheetsProof(self.remote), LocalDrive(self.root/'drive'))
        self.flow.proof.import_history(demo_files(self.root).read_bytes())
        batch = self.flow.process()
        edits = [r['record'].copy() for r in batch['rows']]
        edits[0]['requester'] = '=literal'
        self.flow.approve(batch, edits, approved=True)

    def test_backup_has_logical_base_and_raw_journal(self):
        target = self.root/'copia.xlsx'
        rows = len(self.remote.rows)
        result = backup_sheets(target, lambda: self.flow)
        self.assertEqual(result, {'linhas_journal': rows, 'base_logica': True})
        self.assertEqual(self.remote.commits, 3) # Histórico, lote, arquivamento: a cópia não grava.
        book = load_workbook(target)
        self.assertEqual(book.sheetnames, ['JANUS 2026', 'journal'])
        journal = [[c if c is not None else '' for c in r] for r in book['journal'].iter_rows(values_only=True)]
        self.assertEqual(journal, [r + ['']*(6-len(r)) for r in self.remote.rows])
        expected = BytesIO(); self.flow.export(expected); expected.seek(0)
        self.assertEqual(list(book['JANUS 2026'].values), list(load_workbook(expected)['JANUS 2026'].values))
        self.assertEqual(book['JANUS 2026'].cell(5, 3).data_type, 's')

    def test_refuses_existing_destination(self):
        target = self.root/'copia.xlsx'
        target.write_bytes(b'anterior')
        components = Mock()
        with self.assertRaisesRegex(ValueError, 'já existe'):
            backup_sheets(target, components)
        components.assert_not_called()
        self.assertEqual(target.read_bytes(), b'anterior')

    def test_inconsistent_journal_still_saves_raw_copy(self):
        self.remote.rows.append(['C', 'lote', 'hash', '1', '1', ''])
        result = backup_sheets(self.root/'copia.xlsx', lambda: self.flow)
        self.assertFalse(result['base_logica'])
        self.assertEqual(load_workbook(self.root/'copia.xlsx').sheetnames, ['journal'])


if __name__ == '__main__':
    unittest.main()
