"""Prova append-only de operador único. Nenhum teste aceita perda/duplicação lógica."""
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from datetime import datetime
import hashlib
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook
from janus import HEADERS, Record, extract
from drive_adapter import LocalDrive
from storage import loads
from tests.fixtures import text, pdf
from experiments.sheets_poc.fake import SimulatedSheets
from experiments.sheets_poc.store import SheetsProof, Coordinator, Busy, Unconfirmed


def history_bytes(changed=False, empty=False):
    book=Workbook(); sheet=book.active; sheet.title='JANUS 2026'
    sheet.append([]); sheet.append(HEADERS)
    if not empty:
        for order in (10,20):
            sheet.append([order,datetime(2026,1,2),'Outra pessoa inventada' if changed else 'Pessoa histórica inventada',
                          None,'C000000000000000009','JANUS',None,None,None,None])
        sheet.append([None]*10)
    stream=BytesIO(); book.save(stream); book.close()
    return stream.getvalue()


class SheetsProofTests(unittest.TestCase):
    def setUp(self):
        self.remote=SimulatedSheets(); self.coordinator=Coordinator()
        self.store=SheetsProof(self.remote,self.coordinator)
        self.store.import_history(history_bytes(empty=True))
        self.item=('file1',Record(protocol='C000000000000000001',requester='Pessoa inventada'),hashlib.sha256(b'pdf1').hexdigest())
        self.second=('file2',replace(self.item[1],protocol='C000000000000000002'),self.item[2])

    def fresh(self):
        return SheetsProof(self.remote,Coordinator()) # Nenhuma memória compartilhada do cliente.

    def test_approval_and_invalid_batch_do_not_append(self):
        before=self.remote.get()
        with self.assertRaises(ValueError):self.store.approve([self.item])
        with self.assertRaises(ValueError):
            self.store.approve([self.item,('bad',Record(protocol='?'),'a'*64)],approved=True)
        self.assertEqual(self.remote.get(),before)

    def test_double_click_is_blocked(self):
        entered=threading.Event(); release=threading.Event()
        def hook():
            entered.set()
            if not release.wait(5):raise TimeoutError('Teste travado')
        self.remote.after_get=hook
        with ThreadPoolExecutor(max_workers=1) as pool:
            future=pool.submit(self.store.approve,[self.item],approved=True)
            try:
                self.assertTrue(entered.wait(5))
                with self.assertRaises(Busy):
                    SheetsProof(self.remote,self.coordinator).approve([self.item],approved=True)
            finally:
                self.remote.after_get=None; release.set()
            self.assertEqual(future.result()['inserted'],1)
        self.assertEqual(len(self.store.state()['records']),1)

    def test_identifiers_stable_under_restart_and_input_reordering(self):
        first=self.store.approve([self.item,self.second],approved=True)
        before=self.remote.commits
        repeated=self.fresh().approve([self.second,self.item],approved=True)
        self.assertEqual(first,repeated)
        self.assertEqual(self.remote.commits,before)
        record=self.store.state()['records'][self.item[1].protocol]
        self.assertEqual(record['record_id'],'JANUS:'+self.item[1].protocol)

    def test_reprocessing_normalized_protocol_and_intralot_duplicate(self):
        duplicate=('duplicate',replace(self.item[1],protocol='c-0000.0000.0000.0000/01'),self.item[2])
        first=self.store.approve([self.item,duplicate],approved=True)
        self.assertEqual((first['inserted'],first['ignored']),(1,1))
        result=self.fresh().approve([duplicate],approved=True)
        self.assertEqual((result['inserted'],result['ignored']),(0,1))
        self.assertEqual(len(self.store.state()['records']),1)

    def test_timeout_before_commit_reprocesses_after_cold_restart(self):
        self.remote.failures=['before']
        with self.assertRaises(Unconfirmed):self.store.approve([self.item],approved=True)
        self.assertEqual(len(self.store.state()['records']),0)
        self.assertEqual(self.fresh().approve([self.item],approved=True)['inserted'],1)
        self.assertEqual(len(self.store.state()['records']),1)

    def test_timeout_after_commit_reconciles_without_extra_append(self):
        self.remote.failures=['after']
        first=self.store.approve([self.item],approved=True)
        before=self.remote.commits
        self.assertEqual(self.fresh().approve([self.item],approved=True),first)
        self.assertEqual(self.remote.commits,before)

    def test_lost_response_and_failed_read_then_cold_restart(self):
        get=self.remote.get; initial=self.remote.commits
        def read():
            if self.remote.commits>initial:raise ConnectionError('Leitura indisponível')
            return get()
        self.remote.get=read; self.remote.failures=['after']
        with self.assertRaises(ConnectionError):self.store.approve([self.item],approved=True)
        self.remote.get=get
        self.assertEqual(self.fresh().approve([self.item],approved=True)['inserted'],1)
        self.assertEqual(len(self.store.state()['records']),1)

    def test_cold_restart_in_transit_then_later_batch_no_overwrite(self):
        self.remote.failures=['deferred']
        with self.assertRaises(Unconfirmed):self.store.approve([self.item],approved=True)
        fresh=self.fresh()
        fresh.approve([self.second],approved=True)
        prefix=[r.copy() for r in self.remote.rows]
        self.remote.deliver_deferred()
        self.assertEqual(self.remote.rows[:len(prefix)],prefix)
        self.assertEqual(set(fresh.state()['records']),{self.item[1].protocol,self.second[1].protocol})
        output=BytesIO();fresh.export(output);output.seek(0)
        book=load_workbook(output)
        self.assertEqual([book.active.cell(i,5).value for i in (3,4)],
                         [self.second[1].protocol,self.item[1].protocol])
        self.assertEqual([book.active.cell(i,1).value for i in (3,4)],[1,2])
        book.close()

    def test_repeated_inflight_requests_add_physical_but_no_logical_duplicates(self):
        self.remote.failures=['deferred','deferred']
        for _ in range(2):
            with self.assertRaises(Unconfirmed):self.fresh().approve([self.item],approved=True)
        self.fresh().approve([self.item],approved=True)
        self.remote.deliver_deferred(reverse=True,copies=2)
        self.assertEqual(sum(r[0]=='R' for r in self.remote.rows),5)
        self.assertEqual(len(self.store.state()['records']),1)
        output=BytesIO();self.store.export(output);output.seek(0)
        book=load_workbook(output)
        self.assertEqual(book.active.max_row,3)
        book.close()

    def test_delayed_overlapping_batch_does_not_hide_distinct_records(self):
        third=('file3',replace(self.item[1],protocol='C000000000000000003'),self.item[2])
        self.remote.failures=['deferred']
        with self.assertRaises(Unconfirmed):self.store.approve([self.item,self.second],approved=True)
        self.fresh().approve([self.second,third],approved=True)
        self.remote.deliver_deferred(copies=2)
        self.assertEqual(set(self.store.state()['records']),{i[1].protocol for i in (self.item,self.second,third)})
        self.assertEqual(len(self.store.state()['records']),3)

    def test_same_protocol_different_approved_content_first_confirmed_wins(self):
        self.remote.failures=['deferred']
        with self.assertRaises(Unconfirmed):self.store.approve([self.item],approved=True)
        changed=('other-pdf',replace(self.item[1],requester='Correção sintética confirmada'),self.item[2])
        self.fresh().approve([changed],approved=True)
        self.remote.deliver_deferred()
        record=next(iter(self.store.state()['records'].values()))
        self.assertEqual(record['row'][2],'Correção sintética confirmada')
        drive=Mock();self.store.resume(drive)
        drive.archive.assert_called_once_with('other-pdf',self.item[2])

    def test_stale_read_after_success_never_causes_duplicate_export(self):
        stale=self.remote.get()
        self.remote.read_faults=[stale,stale]
        with self.assertRaises(Unconfirmed):self.store.approve([self.item],approved=True)
        first_rows=len(self.remote.rows)
        self.assertEqual(self.fresh().approve([self.item],approved=True)['inserted'],1)
        self.assertEqual(len(self.remote.rows),first_rows)

    def test_history_retry_deduplicates_import_but_preserves_original_duplicates(self):
        self.remote=SimulatedSheets();self.store=SheetsProof(self.remote,Coordinator())
        source=history_bytes();before=hashlib.sha256(source).hexdigest()
        self.remote.failures=['deferred']
        with self.assertRaises(Unconfirmed):self.store.import_history(source)
        self.fresh().import_history(source)
        self.remote.deliver_deferred(copies=2)
        self.assertEqual(len(self.store.state()['history']),3)
        self.assertEqual(len({r[4] for r in self.store.state()['history'] if r[4]}),1)
        duplicate=('historical',Record(protocol='c-0000.0000.0000.0000/09'),'a'*64)
        item=(self.item[0],replace(self.item[1],requester='=1+1'),self.item[2])
        result=self.store.approve([item,duplicate],approved=True)
        self.assertEqual((result['inserted'],result['ignored']),(1,1))
        with self.assertRaises(ValueError):self.store.import_history(history_bytes(True))
        output=BytesIO();self.store.export(output);output.seek(0)
        book=load_workbook(output);sheet=book['JANUS 2026']
        self.assertEqual(tuple(c.value for c in sheet[2]),HEADERS)
        self.assertEqual(sheet.cell(3,2).value,datetime(2026,1,2))
        self.assertEqual([sheet.cell(i,1).value for i in (3,4,5,6)],[10,20,None,21])
        self.assertEqual(sheet.cell(6,3).value,'=1+1');self.assertEqual(sheet.cell(6,3).data_type,'s')
        book.close();self.assertEqual(hashlib.sha256(source).hexdigest(),before)

    def test_unconfirmed_write_never_archives(self):
        drive=Mock();self.remote.failures=['deferred']
        with self.assertRaises(Unconfirmed):self.store.approve([self.item],approved=True)
        self.assertEqual(self.fresh().resume(drive)['archived'],0)
        drive.archive.assert_not_called()
        self.remote.deliver_deferred()
        self.assertEqual(self.fresh().resume(drive)['archived'],1)
        drive.archive.assert_called_once_with(self.item[0],self.item[2])

    def test_late_replay_never_undoes_archive_acknowledgement(self):
        self.remote.failures=['deferred']
        with self.assertRaises(Unconfirmed):self.store.approve([self.item],approved=True)
        self.fresh().approve([self.item],approved=True)
        drive=Mock();self.store.resume(drive)
        self.remote.deliver_deferred(copies=2)
        self.assertEqual(self.fresh().resume(drive)['archived'],0)
        drive.archive.assert_called_once()

    def test_existing_drive_recovers_moved_pdf_and_lost_ack(self):
        with tempfile.TemporaryDirectory() as tmp:
            drive=LocalDrive(Path(tmp)/'drive')
            path=drive.root/'entrada'/'file1.pdf';pdf(path,text(self.item[1].protocol))
            data=drive.read(path.name);item=(path.name,extract(path),hashlib.sha256(data).hexdigest())
            self.store.approve([item],approved=True)
            archive=drive.archive
            def lost(file_id,digest):
                archive(file_id,digest);raise TimeoutError('Movido; confirmação perdida')
            drive.archive=lost
            self.assertEqual(self.store.resume(drive)['failures'],[path.name])
            drive.archive=archive
            self.remote.failures=['deferred']
            with self.assertRaises(Unconfirmed):self.store.resume(drive)
            self.assertEqual(self.fresh().resume(drive)['archived'],1)
            self.remote.deliver_deferred(copies=2)
            self.assertEqual(self.fresh().resume(drive)['archived'],0)
            self.assertEqual((drive.root/'processados'/path.name).read_bytes(),data)

    def test_invalid_manifest_blocks_export_and_archiving(self):
        self.store.approve([self.item],approved=True)
        self.remote.rows[-1][2]='invalid-hash' # Corrupção modelada: não aceitar em silêncio.
        with self.assertRaises(ValueError):self.store.export(BytesIO())
        drive=Mock()
        with self.assertRaises(ValueError):self.store.resume(drive)
        drive.archive.assert_not_called()

    def test_large_cell_and_file_rebinding_are_rejected(self):
        item=(self.item[0],replace(self.item[1],description='x'*50001),self.item[2])
        before=self.remote.get()
        with self.assertRaises(ValueError):self.store.approve([item],approved=True)
        self.assertEqual(self.remote.get(),before)
        self.store.approve([self.item],approved=True)
        changed=(self.item[0],replace(self.item[1],protocol=self.second[1].protocol),self.item[2])
        with self.assertRaises(ValueError):self.store.approve([changed],approved=True)

    def test_all_arrival_orders_of_three_delayed_batches_and_duplicate_deliveries(self):
        from itertools import permutations
        third=('file3',replace(self.item[1],protocol='C000000000000000003'),self.item[2])
        fourth=('file4',replace(self.item[1],protocol='C000000000000000004'),self.item[2])
        fifth=('file5',replace(self.item[1],protocol='C000000000000000005'),self.item[2])
        source=history_bytes(empty=True)
        for order in permutations(range(3)):
            with self.subTest(arrival_order=order):
                remote=SimulatedSheets();client=SheetsProof(remote)
                client.import_history(source)
                batches=[[self.item,self.second],[self.second,third],[fourth]]
                for items in batches:
                    remote.failures=['deferred']
                    with self.assertRaises(Unconfirmed):SheetsProof(remote).approve(items,approved=True)
                SheetsProof(remote).approve([fifth],approved=True)
                prefix=[row.copy() for row in remote.rows]
                pending=remote.deferred.copy();remote.deferred=[pending[i] for i in order]
                remote.deliver_deferred(copies=2)
                self.assertEqual(remote.rows[:len(prefix)],prefix)
                state=SheetsProof(remote).state()
                self.assertEqual(set(state['records']),{i[1].protocol for i in (self.item,self.second,third,fourth,fifth)})
                self.assertEqual(next(iter(state['records'])),fifth[1].protocol)
                for items in batches:SheetsProof(remote).approve(items,approved=True)
                self.assertEqual(len(SheetsProof(remote).state()['records']),5)
                drive=Mock();self.assertEqual(SheetsProof(remote).resume(drive)['archived'],5)
                self.assertEqual(drive.archive.call_count,5)
