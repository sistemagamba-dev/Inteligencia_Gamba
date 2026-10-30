from dataclasses import replace
from datetime import date
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from janus import (HEADERS, History, Record, TrialStore, creation_time, extract,
                   normalize, parse_text, read_history, valid_protocol)
from tests.fixtures import PROTOCOL, pdf, text


class ExtractionTests(unittest.TestCase):
    def test_normalization(self):
        self.assertEqual(normalize(' c-1234.5678.9012.3456/78 '), PROTOCOL)
        for value in ['', 'C123', 'X123456789012345678', 'C1234567890123456789']:
            self.assertFalse(valid_protocol(value))

    def test_zeladoria_mapping(self):
        r = parse_text(text(), 'D:20261028010000Z')
        self.assertEqual(r.row(10), [10, None, 'Pessoa Sintetica Responsavel',
            '2026-10-27', PROTOCOL, 'JANUS', 'Manutencao sintetica',
            'Rua Exemplo 100', 'Pedido sintetico para teste.', 'Bairro Exemplo'])
        self.assertIn('00000-000', r.full_address)
        self.assertEqual(r.warnings, [])

    def test_event_mapping(self):
        r = parse_text(text(events=True), 'D:20261028010000Z')
        self.assertEqual(r.category, 'cultural')
        self.assertEqual(r.requester, 'Pessoa Sintetica Responsavel')
        self.assertEqual(r.address, 'Rua Exemplo 100 de 100 ao fim - lado par')
        self.assertEqual(r.neighborhood, 'Bairro Exemplo')
        self.assertEqual(r.description, 'Pedido sintetico para teste.')
        self.assertEqual(r.warnings, [])

    def test_pdf_round_trip_both_layouts(self):
        with tempfile.TemporaryDirectory() as tmp:
            for events in [False, True]:
                p = Path(tmp)/'synthetic.pdf'
                pdf(p, text(events=events), pages=4 if events else 2)
                self.assertEqual(extract(p), parse_text(text(events=events), 'D:20261028010000Z'))

    def test_metadata_missing_malformed_and_naive(self):
        for raw in [None, 123, '', 'D:20260230000000Z', 'D:20261028010000', 'invalid']:
            self.assertIsNone(creation_time(raw))
            self.assertIn('ausente_created', parse_text(text(), raw).warnings)

    def test_timezone_midnight_and_offsets(self):
        cases = {'D:20261028010000Z': '2026-10-27T22:00:00-03:00',
                 "D:20261028010000+02'00'": '2026-10-27T20:00:00-03:00',
                 "D:20261028010000-03'00'": '2026-10-28T01:00:00-03:00',
                 'D:20180101010000Z': '2017-12-31T23:00:00-02:00'}
        for raw, expected in cases.items():
            self.assertEqual(creation_time(raw).isoformat(), expected)

    def test_invalid_offset_components_leave_date_missing(self):
        for sign in ['+', '-']:
            for components in ['0060', '0099', '2360', '2400', '9900']:
                for zone in [components, components[:2] + "'" + components[2:] + "'"]:
                    raw = 'D:20261028040000' + sign + zone
                    with self.subTest(raw=raw):
                        self.assertIsNone(creation_time(raw))
                        record = parse_text(text(), raw)
                        self.assertEqual(record.created, '')
                        self.assertIn('ausente_created', record.warnings)

    def test_valid_offset_components_convert_to_sao_paulo(self):
        cases = {'+0000': '2026-10-28T01:00:00-03:00',
                 '-0000': '2026-10-28T01:00:00-03:00',
                 '+0059': '2026-10-28T00:01:00-03:00',
                 '-0059': '2026-10-28T01:59:00-03:00',
                 "+05'30'": '2026-10-27T19:30:00-03:00',
                 "-03'30'": '2026-10-28T04:30:00-03:00',
                 "+23'59'": '2026-10-27T01:01:00-03:00',
                 "-23'59'": '2026-10-29T00:59:00-03:00'}
        for zone, expected in cases.items():
            raw = 'D:20261028040000' + zone
            with self.subTest(raw=raw):
                converted = creation_time(raw)
                self.assertEqual(converted.tzinfo.key, 'America/Sao_Paulo')
                self.assertEqual(converted.isoformat(), expected)
                record = parse_text(text(), raw)
                self.assertEqual(record.created, expected[:10])
                self.assertNotIn('ausente_created', record.warnings)

    def test_optional_missing(self):
        r = parse_text(text().replace('Enquadramento: Manutencao sintetica\n', ''), None)
        self.assertIn('ausente_category', r.warnings)
        self.assertTrue(valid_protocol(r.protocol))

    def test_ambiguous_address_no_guess(self):
        r = parse_text(text().replace('Rua Exemplo 100, Bairro Exemplo, Santo André - SP ,\n00000-000',
                                     'Endereco sem bairro'), None)
        self.assertEqual(r.address, 'Endereco sem bairro')
        self.assertEqual(r.neighborhood, '')
        self.assertIn('bairro_ambiguo_revisar', r.warnings)

    def test_conflicting_protocol_blocks(self):
        r = parse_text(text() + '\nProtocolo: C999999999999999999', None)
        self.assertEqual(r.protocol, '')

    def test_unknown_type_reported(self):
        r = parse_text(text().replace('Zeladoria', 'Outro'), None)
        self.assertIn('tipo_nao_suportado_revisar', r.warnings)


class PersistenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.h = History([[7, None, 'Historico sintetico', None,
                           PROTOCOL, 'JANUS', None, None, None, None]] * 2,
                         {PROTOCOL}, 8, 2)
        self.store = TrialStore(self.root/'trial.sqlite', self.h)
        self.record = Record(protocol='C999999999999999999', requester='Pessoa sintetica')
        self.moved = set()

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def count(self):
        return self.store.db.execute('SELECT count(*) FROM records').fetchone()[0]

    def test_human_approval_required(self):
        with self.assertRaises(ValueError):
            self.store.approve([('one', self.record)], approved=False, archive=self.moved.add)
        self.assertEqual(self.count(), 0)
        self.assertEqual(self.moved, set())

    def test_invalid_batch_entirely_blocked(self):
        with self.assertRaises(ValueError):
            self.store.approve([('one', self.record), ('two', Record(protocol='?'))],
                               approved=True, archive=self.moved.add)
        self.assertEqual(self.count(), 0)

    def test_historical_duplicate_and_variants_preserved(self):
        result = self.store.approve([('one', Record(protocol='c-1234.5678.9012.3456/78'))],
                                    approved=True, archive=self.moved.add)
        self.assertEqual(result, {'inserted': 0, 'ignored': 1})
        self.assertEqual(len(self.h.rows), 2)

    def test_duplicate_same_batch_and_repeat(self):
        items = [('one', self.record), ('two', replace(self.record, protocol='c-9999.9999.9999.9999/99'))]
        self.assertEqual(self.store.approve(items, approved=True, archive=self.moved.add),
                         {'inserted': 1, 'ignored': 1})
        self.assertEqual(self.store.approve(items, approved=True, archive=self.moved.add),
                         {'inserted': 0, 'ignored': 2})
        self.assertEqual(self.count(), 1)

    def test_missing_optional_can_be_approved(self):
        r = parse_text(text().replace('Enquadramento: Manutencao sintetica\n', ''), None)
        r.protocol = self.record.protocol
        self.assertEqual(self.store.approve([('one', r)], approved=True, archive=self.moved.add)['inserted'], 1)

    def test_commit_before_archive_failure_restart(self):
        def fail():
            raise OSError('falha sintetica apos commit')
        with self.assertRaises(OSError):
            self.store.approve([('one', self.record)], approved=True,
                               archive=self.moved.add, after_commit=fail)
        self.assertEqual(self.count(), 1)
        self.assertEqual(self.moved, set())
        self.store.close()
        self.store = TrialStore(self.root/'trial.sqlite', self.h)
        self.store.resume(self.moved.add)
        self.assertEqual(self.moved, {'one'})
        self.assertEqual(self.count(), 1)

    def test_move_then_failure_retried_idempotently(self):
        def move_then_fail(file_id):
            self.moved.add(file_id)
            raise OSError('confirmacao perdida')
        with self.assertRaises(OSError):
            self.store.approve([('one', self.record)], approved=True, archive=move_then_fail)
        self.store.resume(self.moved.add)
        self.assertEqual(self.moved, {'one'})
        self.assertEqual(self.count(), 1)
        self.assertEqual(self.store.db.execute('SELECT archived FROM outbox').fetchone()[0], 1)

    def test_transaction_rollback_does_not_archive(self):
        self.store.approve([('one', self.record)], approved=True, archive=self.moved.add)
        with self.assertRaises(ValueError):
            self.store.approve([('two', Record(protocol='C222222222222222222')),
                                ('one', Record(protocol='C333333333333333333'))],
                               approved=True, archive=self.moved.add)
        self.assertEqual(self.count(), 1)
        self.assertEqual(self.moved, {'one'})

    def test_export_preserves_duplicates_and_literal_formula(self):
        self.store.approve([('one', replace(self.record, requester='=1+1'))],
                            approved=True, archive=self.moved.add)
        path = self.root/'export.xlsx'
        self.store.export(path)
        w = load_workbook(path, data_only=False)
        rows = list(w.active.iter_rows(min_row=3, values_only=True))
        self.assertEqual([list(r) for r in rows[:2]], self.h.rows)
        self.assertEqual(rows[2][0], 8)
        self.assertEqual(rows[2][2], '=1+1')
        self.assertEqual(w.active.cell(5, 3).data_type, 's')
        w.close()

    def test_two_connections_no_overwrite(self):
        other = TrialStore(self.root/'trial.sqlite', self.h)
        try:
            self.store.approve([('one', self.record)], approved=True, archive=self.moved.add)
            result = other.approve([('two', replace(self.record, requester='Outra pessoa'))],
                                   approved=True, archive=self.moved.add)
            self.assertEqual(result['ignored'], 1)
            payload = other.db.execute('SELECT payload FROM records').fetchone()[0]
            self.assertNotIn('Outra pessoa', payload)
        finally:
            other.close()


class HistoryTests(unittest.TestCase):
    def test_read_only_history_and_sequence(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'history.xlsx'
            w = Workbook()
            w.active.title = 'JANUS 2026'
            w.active.append([])
            w.active.append(HEADERS)
            for n in [2, 9]:
                w.active.append([n, None, 'Pessoa sintetica', None, PROTOCOL])
            w.save(path)
            w.close()
            before = path.read_bytes()
            h = read_history(path)
            self.assertEqual((len(h.rows), len(h.keys), h.protocol_rows, h.next_order), (2, 1, 2, 10))
            self.assertEqual(path.read_bytes(), before)

    def test_wrong_headers_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)/'history.xlsx'
            w = Workbook()
            w.active.title = 'JANUS 2026'
            w.active.append(['wrong'])
            w.save(p)
            w.close()
            with self.assertRaises(ValueError):
                read_history(p)


if __name__ == '__main__':
    unittest.main()
