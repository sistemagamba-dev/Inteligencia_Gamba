from dataclasses import asdict
from datetime import datetime
import hashlib
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from openpyxl import load_workbook
from drive_adapter import GoogleDrive, LocalDrive
from janus import HEADERS
from manage import demo, restore
from storage import Store, Busy
from workflow import Workflow
from tests.fixtures import pdf, text


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        demo(self.root)
        self.drive = LocalDrive(self.root/'drive')
        self.store = Store(self.root/'janus.sqlite')
        self.flow = Workflow(self.store, self.drive)

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def count(self):
        return self.store.db.execute('SELECT count(*) FROM records').fetchone()[0]

    def batch(self):
        batch = self.flow.process()
        return batch, [r['record'].copy() for r in batch['rows']]

    def test_end_to_end_history_optional_edit_export_restart(self):
        original = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.drive.root/'entrada').iterdir()}
        batch, edits = self.batch()
        self.assertEqual((batch['found'],len(edits),len(batch['duplicates']),len(batch['failures'])), (5,3,1,1))
        self.assertEqual(self.count(), 0)
        edits[0]['requester'] = '=literal corrigido'
        result = self.flow.approve(batch['id'], edits, approved=True)
        self.assertEqual((result['inserted'],result['ignored'],result['failures'],result['archived']), (3,1,1,3))
        self.store.close()
        self.store = Store(self.root/'janus.sqlite')
        self.flow = Workflow(self.store, self.drive)
        output = BytesIO()
        self.store.export(output)
        output.seek(0)
        book = load_workbook(output)
        sheet = book['JANUS 2026']
        self.assertEqual(tuple(c.value for c in sheet[2]), HEADERS)
        self.assertEqual([sheet.cell(i,1).value for i in range(3,8)], [4,9,10,11,12])
        self.assertEqual(sheet.cell(5,3).value, '=literal corrigido')
        self.assertEqual(sheet.cell(5,3).data_type, 's')
        self.assertIsNone(sheet.cell(7,4).value)
        book.close()
        self.assertEqual(self.flow.approve(batch['id'], [], approved=True)['inserted'],3) # replay receipt, no write
        self.assertEqual(self.count(),3)
        again = self.flow.process()
        self.assertEqual((again['found'],len(again['rows'])), (2,0))
        for name, digest in original.items():
            paths = [folder/name for folder in (self.drive.root/'entrada',self.drive.root/'processados')]
            self.assertEqual(sum(p.exists() for p in paths),1)
            self.assertEqual(hashlib.sha256(next(p for p in paths if p.exists()).read_bytes()).hexdigest(),digest)

    def test_approval_invalid_protocol_date_and_no_partial_write(self):
        batch, edits = self.batch()
        with self.assertRaises(ValueError):
            self.flow.approve(batch['id'],edits)
        edits[-1]['protocol'] = '?'
        with self.assertRaises(ValueError):
            self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual(self.count(),0)
        edits[-1]['protocol'] = 'C000000000000000003'
        edits[-1]['created'] = '2026-02-30'
        with self.assertRaises(ValueError):
            self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual(self.count(),0)
        self.assertFalse(list((self.drive.root/'processados').iterdir()))

    def test_after_commit_failure_and_restart(self):
        batch, edits = self.batch()
        def fail():
            raise OSError('injetado')
        with self.assertRaises(OSError):
            self.flow.approve(batch['id'],edits,approved=True,after_commit=fail)
        self.assertEqual(self.count(),3)
        self.store.close()
        self.store = Store(self.root/'janus.sqlite')
        self.flow = Workflow(self.store,self.drive)
        self.assertEqual(self.flow.resume()['archived'],3)
        self.assertEqual(self.flow.approve(batch['id'],edits,approved=True)['archived'],0)
        self.assertEqual(self.count(),3)

    def test_one_archive_failure_does_not_stop_other_files(self):
        batch, edits = self.batch()
        archive = self.drive.archive
        def fail_one(file_id, digest):
            if file_id == batch['rows'][0]['file_id']:
                raise OSError('injetado')
            archive(file_id,digest)
        self.drive.archive = fail_one
        result = self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual((result['archived'],len(result['archive_failures'])), (2,1))
        self.drive.archive = archive
        self.assertEqual(self.flow.resume()['archived'],1)

    def test_move_confirmation_lost_is_idempotent(self):
        batch, edits = self.batch()
        archive = self.drive.archive
        def move_and_fail(file_id,digest):
            archive(file_id,digest)
            raise OSError('confirmação perdida')
        self.drive.archive = move_and_fail
        result = self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual(len(result['archive_failures']),3)
        self.drive.archive = archive
        self.assertEqual(self.flow.resume()['archived'],3)
        self.assertEqual(self.count(),3)

    def test_modified_pdf_before_approval_or_after_commit_preserved(self):
        batch, edits = self.batch()
        source = self.drive.root/'entrada'/batch['rows'][0]['file_id']
        source.write_bytes(b'changed')
        with self.assertRaises(ValueError):
            self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual(self.count(),0)
        # Another valid batch, mutate after confirmed persistence.
        batch, edits = self.batch()
        source = self.drive.root/'entrada'/batch['rows'][0]['file_id']
        result = self.flow.approve(batch['id'],edits,approved=True,after_commit=lambda:source.write_bytes(b'changed'))
        self.assertEqual(len(result['archive_failures']),1)
        self.assertTrue(source.exists())

    def test_concurrent_previews_never_duplicate_or_archive_rejected_files(self):
        batch, edits = self.batch()
        other = Store(self.root/'janus.sqlite')
        try:
            with self.store.exclusive():
                with self.assertRaises(Busy):
                    Workflow(other,self.drive).process()
            # A stale preview whose records were inserted by another batch remains safe.
            another, other_edits = self.batch()
            self.flow.approve(batch['id'],edits,approved=True)
            with self.assertRaises(FileNotFoundError):
                Workflow(other,self.drive).approve(another['id'],other_edits,approved=True)
            self.assertEqual(self.count(),3)
        finally:
            other.close()

    def test_duplicate_after_manual_edit_remains_input(self):
        batch, edits = self.batch()
        edits[1]['protocol'] = edits[0]['protocol']
        result = self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual((result['inserted'], result['ignored']), (2,2))
        self.assertTrue((self.drive.root/'entrada'/batch['rows'][1]['file_id']).exists())

    def test_kernel_lock_between_independent_processes(self):
        import subprocess
        import sys
        child = subprocess.Popen([sys.executable, '-c',
            "from storage import Store; import sys; s=Store(sys.argv[1]); "
            "ctx=s.exclusive(); ctx.__enter__(); print('locked',flush=True); "
            "sys.stdin.read(1); ctx.__exit__(None,None,None); s.close()",
            str(self.store.path)], stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
        try:
            self.assertEqual(child.stdout.readline().strip(), 'locked')
            with self.assertRaises(Busy):
                self.flow.process()
            self.assertEqual(self.count(), 0)
        finally:
            child.communicate('x', timeout=10)
        self.assertEqual(child.returncode, 0)

    def test_persistent_history_preserves_datetime_cells_and_rejects_replacement(self):
        from openpyxl import Workbook
        root = self.root/'dates'
        root.mkdir()
        source = root/'dates.xlsx'
        book = Workbook()
        book.active.title = 'JANUS 2026'
        book.active.append([])
        book.active.append(HEADERS)
        book.active.append([22, datetime(2026,1,2), 'Histórico sintético', datetime(2026,1,1), 'C111111111111111111'])
        book.save(source)
        book.close()
        store = Store(root/'dates.sqlite')
        store.import_history(source)
        store.close()
        store = Store(root/'dates.sqlite')
        try:
            store.export(root/'dates-export.xlsx')
            exported = load_workbook(root/'dates-export.xlsx')
            self.assertEqual(exported.active.cell(3,2).value, datetime(2026,1,2))
            exported.close()
            with self.assertRaises(ValueError):
                store.import_history(self.root/'historico-sintetico.xlsx')
        finally:
            store.close()

    def test_history_date_duplicate_backup_restore_and_refusal_replace(self):
        # Existing demo history includes intentional duplicates, persisted without source.
        history = self.root/'historico-sintetico.xlsx'
        self.assertFalse(self.store.import_history(history))
        history.unlink()
        backup = self.root/'safe.bak'
        self.store.backup(backup)
        batch, edits = self.batch()
        self.flow.approve(batch['id'],edits,approved=True)
        self.assertEqual(self.count(),3)
        self.store.close()
        restore(self.root/'janus.sqlite',backup)
        self.store = Store(self.root/'janus.sqlite')
        self.assertEqual(self.count(),0)
        self.assertEqual(len(self.store.history.rows),2)
        self.assertEqual(len(list(self.root.glob('*.before-restore-*.bak'))),1)
        before = (self.root/'janus.sqlite').read_bytes()
        with self.assertRaises(sqlite3.DatabaseError):
            broken = self.root/'broken.bak'; broken.write_bytes(b'broken')
            restore(self.root/'janus.sqlite',broken)
        self.assertEqual(before,(self.root/'janus.sqlite').read_bytes())


import sqlite3

class GoogleContractTests(unittest.TestCase):
    def response(self, data=None, content=b''):
        r = Mock(); r.json.return_value = data; r.content = content
        return r

    def adapter(self):
        self.session = Mock()
        root = {'name':'JANUS_TESTE sintetico','mimeType':'application/vnd.google-apps.folder'}
        child = {'parents':['root'],'mimeType':'application/vnd.google-apps.folder'}
        self.session.request.side_effect = [self.response(root)] + [self.response(child)]*3
        return GoogleDrive(self.session,'root','inbox','done','failed')

    def test_pages_and_parent_scoped_download(self):
        drive = self.adapter()
        self.session.request.side_effect = [self.response({'files':[{'id':'one'}],'nextPageToken':'next'}),
                                            self.response({'files':[{'id':'two'}]})]
        self.assertEqual([f.id for f in drive.pending()],['one','two'])
        self.assertEqual(self.session.request.call_args.kwargs['params']['pageToken'],'next')
        self.session.request.side_effect = [self.response({'parents':['outside'],'mimeType':'application/pdf'})]
        with self.assertRaises(ValueError): drive.read('one')

    def test_archive_verify_after_move_and_retry_target(self):
        drive = self.adapter()
        data = b'synthetic'
        digest = hashlib.sha256(data).hexdigest()
        source = {'parents':['inbox']}; target = {'parents':['done']}
        self.session.request.side_effect = [self.response(source), self.response(content=data),
                                            self.response({}), self.response(target)]
        drive.archive('one',digest)
        request = self.session.request.call_args_list[-2]
        self.assertEqual(request.args[0],'PATCH')
        self.assertEqual(request.kwargs['params']['removeParents'],'inbox')
        self.assertEqual(request.kwargs['params']['addParents'],'done')
        self.session.request.side_effect = [self.response(target),self.response(content=data)]
        drive.archive('one',digest)
        self.assertEqual(self.session.request.call_args.args[0],'GET')

    def test_modified_content_permission_failure_and_wrong_root(self):
        drive = self.adapter()
        self.session.request.side_effect = [self.response({'parents':['inbox']}),self.response(content=b'changed')]
        with self.assertRaises(ValueError): drive.archive('one','wrong-digest')
        response = self.response(); response.raise_for_status.side_effect = OSError('403')
        self.session.request.side_effect = [response]
        with self.assertRaises(OSError): drive.pending()
        self.session.request.side_effect = [self.response({'name':'Operacional'})]
        with self.assertRaises(ValueError): GoogleDrive(self.session,'root','inbox','done','failed')
