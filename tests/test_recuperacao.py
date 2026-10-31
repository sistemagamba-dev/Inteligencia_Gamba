"""Ensaio de recuperação: importar histórico, registrar, copiar, perder a base e
recuperar numa planilha nova só com os comandos existentes (backup-sheets e
import-history-sheets). Dados sintéticos, Sheets simulado, Drive local."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import shutil
import tempfile
from types import SimpleNamespace
import unittest

from openpyxl import Workbook, load_workbook
from drive_adapter import LocalDrive
from experiments.sheets_poc.fake import SimulatedSheets
from experiments.sheets_poc.store import SheetsProof, Coordinator
from janus import HEADERS, normalize
from manage import backup_sheets, import_history_sheets
from tests.fixtures import pdf, text
from workflow import SheetsWorkflow

HIST = 'C000000000000000009'


def history_file(path):
    book = Workbook(); sheet = book.active; sheet.title = 'JANUS 2026'
    sheet.append([]); sheet.append(HEADERS)
    sheet.append([1, None, 'Pessoa histórica A', datetime(2025, 1, 2), HIST, 'JANUS', 'Cat', 'Rua 1', 'Pedido', 'Bairro'])
    sheet.append([2, None, 'Pessoa histórica B', datetime(2025, 3, 4), 'c-0000.0000.0000.0000/09', 'JANUS', None, None, None, None])
    sheet.append([None] * 10)  # Linha vazia no meio do histórico.
    sheet.append([None, 'texto', '=literal', None, 'SEM PROTOCOLO', 'JANUS', None, None, None, None])
    sheet.append([3, None, 'Pessoa histórica C', datetime(2026, 10, 1), 'C000000000000000010', 'JANUS', 'Cat', 'Rua 3', 'Outro', 'Bairro'])
    sheet['C6'].data_type = 's'  # Texto literal iniciado por "=", não fórmula.
    book.save(path); book.close()


def rows(proof):
    out = BytesIO(); proof.export(out); out.seek(0)
    book = load_workbook(out, read_only=True)
    values = [list(r) for r in book['JANUS 2026'].iter_rows(min_row=3, max_col=10, values_only=True)]
    book.close()
    return values


class RecoveryRehearsal(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp()); self.addCleanup(shutil.rmtree, self.tmp)
        self.drive = LocalDrive(self.tmp/'drive')
        history_file(self.tmp/'historico.xlsx')

    def base(self):
        remote = SimulatedSheets()
        proof = SheetsProof(remote, Coordinator())
        return remote, proof, (lambda: SimpleNamespace(proof=proof))

    def add_pdf(self, name, protocol):
        pdf(self.drive.root/'entrada'/f'{name}.pdf', text(protocol))

    def register(self, proof):
        flow = SheetsWorkflow(proof, self.drive)
        batch = flow.process()
        return batch, flow.approve(batch, [r['record'] for r in batch['rows']], approved=True)

    def test_backup_loss_and_recovery_in_new_base(self):
        # 1. Base A: importação única do histórico e registro de novos protocolos.
        _, proof_a, components_a = self.base()
        self.assertTrue(import_history_sheets(self.tmp/'historico.xlsx', components_a))
        for i in (1, 2, 3):
            self.add_pdf(f'novo{i}', f'C77030000000000000{i}')
        self.add_pdf('dup', HIST)
        batch, result = self.register(proof_a)
        self.assertEqual((result['inserted'], result['ignored'], result['archived']), (3, 1, 3))
        before = rows(proof_a)
        # 2. Cópia de segurança; 3. perda da base A.
        summary = backup_sheets(self.tmp/'backup.xlsx', components_a)
        self.assertTrue(summary['base_logica'])
        del proof_a, components_a
        # 4. Recuperação: importar a aba JANUS 2026 da cópia numa planilha nova (base B).
        remote_b, proof_b, components_b = self.base()
        self.assertTrue(import_history_sheets(self.tmp/'backup.xlsx', components_b))
        after = rows(proof_b)
        self.assertEqual(after, before)  # Mesmos valores A:J, na mesma ordem.
        protocols = [r[4] for r in after if r[4]]
        self.assertEqual(sum(normalize(p) == HIST for p in protocols), 2)  # Duplicata histórica preservada.
        self.assertEqual(len(protocols), 7)  # 4 históricos com protocolo + 3 registrados.
        self.assertIn([None] * 10, after)  # Linha vazia intermediária preservada.
        self.assertIn('=literal', [r[2] for r in after])
        # 5. Operação continua na base B sem duplicar nem perder.
        self.add_pdf('novo4', 'C770300000000000004')
        batch, result = self.register(proof_b)
        self.assertEqual((result['inserted'], result['ignored']), (1, 1))  # novo4 entra; dup segue ignorado.
        final = rows(proof_b)
        self.assertEqual(final[:-1], before)
        self.assertEqual(final[-1][0], max(r[0] for r in before if isinstance(r[0], int)) + 1)
        # Repetir a importação da cópia na base B não regrava nada.
        appends = remote_b.commits
        self.assertFalse(import_history_sheets(self.tmp/'backup.xlsx', components_b))
        self.assertEqual(remote_b.commits, appends)

    def test_records_after_last_backup_are_recovered_by_reprocessing(self):
        _, proof_a, components_a = self.base()
        import_history_sheets(self.tmp/'historico.xlsx', components_a)
        self.add_pdf('antes', 'C770300000000000011')
        self.register(proof_a)
        backup_sheets(self.tmp/'backup.xlsx', components_a)
        self.add_pdf('depois', 'C770300000000000012')  # Registrado depois da cópia e perdido com a base A.
        self.register(proof_a)
        _, proof_b, components_b = self.base()
        import_history_sheets(self.tmp/'backup.xlsx', components_b)
        # Procedimento: devolver à entrada os PDFs de processados posteriores à cópia (ou todos).
        for path in (self.drive.root/'processados').iterdir():
            path.rename(self.drive.root/'entrada'/path.name)
        batch, result = self.register(proof_b)
        self.assertEqual((result['inserted'], result['ignored']), (1, 1))  # Só o perdido entra; o outro é duplicado.
        keys = [normalize(r[4]) for r in rows(proof_b) if r[4]]
        self.assertEqual(keys.count('C770300000000000011'), 1)
        self.assertEqual(keys.count('C770300000000000012'), 1)

    def test_backup_refuses_overwrite_and_import_refuses_different_history(self):
        _, proof_a, components_a = self.base()
        import_history_sheets(self.tmp/'historico.xlsx', components_a)
        backup_sheets(self.tmp/'backup.xlsx', components_a)
        with self.assertRaises(ValueError):
            backup_sheets(self.tmp/'backup.xlsx', components_a)
        # A base de origem nunca recebe a cópia como "outro histórico".
        with self.assertRaisesRegex(ValueError, 'Não substituir'):
            import_history_sheets(self.tmp/'backup.xlsx', components_a)


if __name__ == '__main__':
    unittest.main()
