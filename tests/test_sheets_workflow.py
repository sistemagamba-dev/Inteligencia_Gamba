"""SheetsWorkflow sobre o diário append-only simulado. Somente dados sintéticos."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from openpyxl import load_workbook
from drive_adapter import LocalDrive
from experiments.sheets_poc.fake import SimulatedSheets
from experiments.sheets_poc.store import SheetsProof, Coordinator, Unconfirmed
from manage import demo_files
from sheets_adapter import GoogleSheets
from workflow import SheetsWorkflow


class SheetsWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.history = demo_files(self.root)
        self.remote = SimulatedSheets()
        self.drive = LocalDrive(self.root/'drive')
        self.flow = self.fresh()
        self.flow.proof.import_history(self.history.read_bytes())

    def tearDown(self):
        self.tmp.cleanup()

    def fresh(self):
        return SheetsWorkflow(SheetsProof(self.remote, Coordinator()), self.drive) # Reinício sem memória.

    def batch(self):
        batch = self.flow.process()
        return batch, [r['record'].copy() for r in batch['rows']]

    def folder(self, name):
        return sorted(p.name for p in (self.drive.root/name).iterdir())

    def test_end_to_end_restart_history_and_export(self):
        batch, edits = self.batch()
        self.assertEqual((batch['found'], len(edits), len(batch['duplicates']), len(batch['failures'])), (5,3,1,1))
        self.assertEqual(self.remote.commits, 1) # Só o histórico: a prévia não grava.
        edits[0]['requester'] = '=literal corrigido'
        result = self.flow.approve(batch, edits, approved=True)
        self.assertEqual((result['inserted'], result['ignored'], result['failures'], result['archived']), (3,1,1,3))
        self.assertEqual(self.folder('entrada'), ['04-duplicado.pdf', '05-invalido.pdf'])
        flow = self.fresh()
        self.assertEqual(flow.process()['rows'], [])
        output = BytesIO(); flow.export(output); output.seek(0)
        sheet = load_workbook(output)['JANUS 2026']
        self.assertEqual([sheet.cell(i,1).value for i in range(3,8)], [4,9,10,11,12])
        self.assertEqual(sheet.cell(5,3).value, '=literal corrigido')
        self.assertEqual(sheet.cell(5,3).data_type, 's')

    def test_requires_approval_valid_fields_and_writes_nothing_partial(self):
        batch, edits = self.batch()
        commits = self.remote.commits
        with self.assertRaises(ValueError):
            self.flow.approve(batch, edits)
        bad = [dict(e) for e in edits]; bad[1]['protocol'] = 'sem protocolo'
        with self.assertRaisesRegex(ValueError, 'protocolo'):
            self.flow.approve(batch, bad, approved=True)
        bad = [dict(e) for e in edits]; bad[2]['created'] = '30/10/2026'
        with self.assertRaisesRegex(ValueError, 'data'):
            self.flow.approve(batch, bad, approved=True)
        with self.assertRaisesRegex(ValueError, 'tabela'):
            self.flow.approve(batch, edits[:1], approved=True)
        self.assertEqual(self.remote.commits, commits)
        self.assertEqual(len(self.folder('processados')), 0)

    def test_lost_response_and_timeout_never_duplicate_or_archive_early(self):
        batch, edits = self.batch()
        self.remote.failures.append('before')
        with self.assertRaises(Unconfirmed):
            self.flow.approve(batch, edits, approved=True)
        self.assertEqual(len(self.folder('processados')), 0)
        self.remote.failures.append('after')
        self.assertEqual(self.fresh().approve(batch, edits, approved=True)['inserted'], 3)
        rows = len(self.remote.rows)
        self.assertEqual(self.fresh().approve(batch, edits, approved=True)['inserted'], 3)
        self.assertEqual(len(self.remote.rows), rows) # Repetição reconhecida pelo ID do lote.
        self.assertEqual(len(self.flow.proof.state()['records']), 3)

    def test_archive_failure_survives_restart_and_resumes(self):
        archive, calls = self.drive.archive, []
        def flaky(file_id, digest):
            calls.append(file_id)
            if len(calls) == 1:
                raise ConnectionError('Drive indisponível')
            return archive(file_id, digest)
        self.drive.archive = flaky
        batch, edits = self.batch()
        result = self.flow.approve(batch, edits, approved=True)
        self.assertEqual((result['archived'], len(result['archive_failures'])), (2,1))
        self.assertIn(calls[0], self.folder('entrada'))
        self.drive.archive = archive
        flow = self.fresh()
        self.assertEqual(len(flow.pending_archives()), 1)
        self.assertEqual(flow.process()['recovery'], {'archived':1, 'archive_failures':[]})
        self.assertEqual(len(self.folder('processados')), 3)

    def test_changed_pdf_and_duplicate_created_by_edit_stay_in_entrada(self):
        batch, edits = self.batch()
        edits[1]['protocol'] = edits[0]['protocol']
        result = self.flow.approve(batch, edits, approved=True)
        self.assertEqual((result['inserted'], result['ignored']), (2,2))
        self.assertIn(batch['rows'][1]['file_id'], self.folder('entrada'))
        other = self.root/'other'
        flow = SheetsWorkflow(SheetsProof(SimulatedSheets()), LocalDrive(other/'drive'))
        flow.proof.import_history(demo_files(other).read_bytes())
        batch = flow.process()
        edits = [r['record'] for r in batch['rows']]
        path = other/'drive'/'entrada'/batch['rows'][0]['file_id']
        path.write_bytes(path.read_bytes() + b'%')
        with self.assertRaisesRegex(ValueError, 'PDF mudou'):
            flow.approve(batch, edits, approved=True)

    def test_base_without_history_is_refused(self):
        flow = SheetsWorkflow(SheetsProof(SimulatedSheets()), self.drive)
        with self.assertRaisesRegex(ValueError, 'equipe técnica'):
            flow.process()


class GoogleSheetsContractTests(unittest.TestCase):
    def test_only_get_and_append_raw_insert_rows_on_journal(self):
        session = Mock()
        sheets = GoogleSheets(session, 'abc_-1')
        sheets.get()
        session.get.assert_called_once_with('https://sheets.googleapis.com/v4/spreadsheets/abc_-1/values/journal!A:F', timeout=60)
        body = {'values':[['R']+['']*5]}
        sheets.append(body, range='journal!A:F', valueInputOption='RAW', insertDataOption='INSERT_ROWS')
        session.post.assert_called_once_with(
            'https://sheets.googleapis.com/v4/spreadsheets/abc_-1/values/journal!A:F:append', timeout=60, json=body,
            params={'valueInputOption':'RAW', 'insertDataOption':'INSERT_ROWS'})
        with self.assertRaises(ValueError):
            sheets.append(body, range='journal!A:F', valueInputOption='USER_ENTERED', insertDataOption='INSERT_ROWS')
        with self.assertRaises(ValueError):
            GoogleSheets(session, '../outra')

    def test_http_errors_propagate(self):
        session = Mock()
        session.get.return_value.raise_for_status.side_effect = RuntimeError('403')
        with self.assertRaises(RuntimeError):
            GoogleSheets(session, 'abc').get()
