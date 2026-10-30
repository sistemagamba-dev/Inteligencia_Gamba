"""Lote sintético com o fluxo da Etapa 2; não usa rede nem dados reais."""
import json
from pathlib import Path
import tempfile
from time import perf_counter
from openpyxl import load_workbook
from tests.fixtures import pdf, text
from drive_adapter import LocalDrive
from storage import Store
from workflow import Workflow


def run():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        drive = LocalDrive(root/'drive')
        store = Store(root/'test.sqlite')
        try:
            for n in range(100):
                pdf(drive.root/'entrada'/f'{n:03}.pdf',text(f'C{n:018}',events=bool(n%2)),pages=4 if n%2 else 2)
            start = perf_counter()
            flow = Workflow(store,drive)
            batch = flow.process()
            preview = perf_counter()
            assert store.db.execute('SELECT count(*) FROM records').fetchone()[0] == 0
            result = flow.approve(batch['id'],[r['record'] for r in batch['rows']],approved=True)
            approved = perf_counter()
            repeated = flow.process()
            finished = perf_counter()
            assert repeated['found'] == 0
            assert result['inserted'] == result['archived'] == 100
            store.export(root/'export.xlsx')
            book = load_workbook(root/'export.xlsx')
            assert book.active.max_row == 102
            book.close()
            assert len(list((drive.root/'processados').glob('*.pdf'))) == 100
            return {'data':'100 PDFs inteiramente sintéticos / 300 páginas / 50 de cada layout',
                    'preview_rows':len(batch['rows']),'records_before_approval':0,'result':result,
                    'repeat_found':repeated['found'],'xlsx_data_rows':100,
                    'seconds':{'scan_and_extract':round(preview-start,6),
                               'commit_and_local_archive':round(approved-preview,6),
                               'repeat':round(finished-approved,6)},
                    'google_network_included':False}
        finally:
            store.close()


if __name__ == '__main__':
    print(json.dumps(run(),indent=2))
